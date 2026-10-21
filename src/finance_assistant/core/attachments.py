"""Read user-uploaded files (PDF, images, Word, Excel, CSV, text) into plain text.

Each file becomes a document dict that agents can quote from:
    {"name": "statement.pdf", "kind": "pdf", "text": "...", "note": "..."}

Images (and scanned PDFs with no text layer) are transcribed by the configured
LLM's vision support. In offline mock mode they are kept with a note instead.
Only the passages most relevant to the question are sent to the LLM, so large
files don't blow up the prompt.
"""
from __future__ import annotations

import io
import json
import logging
import re
from pathlib import Path

from ..rag.embeddings import tokenize

log = logging.getLogger(__name__)

TEXT_TYPES = {"txt", "md", "csv", "json"}
IMAGE_TYPES = {"png", "jpg", "jpeg", "webp", "gif"}
OFFICE_TYPES = {"pdf", "docx", "xlsx"}
SUPPORTED_TYPES = sorted(TEXT_TYPES | IMAGE_TYPES | OFFICE_TYPES)
MAX_TEXT_CHARS = 60_000    # kept per file
MAX_IMAGE_SIDE = 1600      # pixels; larger images are shrunk before sending to the LLM
MAX_SCANNED_PAGES = 3      # page images transcribed for PDFs without a text layer
IMAGE_PROMPT = ("Transcribe all text and numbers in this image exactly (tables as rows). Then add one line "
                "starting 'Description:' saying what the image shows. Do not give advice.")


def read_file(name: str, data: bytes, llm=None) -> dict:
    """Extract text from one uploaded file. Never raises: problems go in `note`."""
    kind = Path(name).suffix.lower().lstrip(".")
    doc = {"name": name, "kind": kind, "text": "", "note": ""}
    try:
        if kind == "pdf":
            doc["text"], doc["note"] = _read_pdf(data, llm)
        elif kind in IMAGE_TYPES:
            doc["text"], doc["note"] = _read_image(data, llm)
        elif kind == "docx":
            doc["text"] = _read_docx(data)
        elif kind == "xlsx":
            doc["text"] = _read_xlsx(data)
        elif kind == "json":
            doc["text"] = json.dumps(json.loads(_decode(data)), indent=1)
        elif kind in TEXT_TYPES:
            doc["text"] = _decode(data)
        else:
            doc["note"] = f"File type .{kind} isn't supported. Try PDF, an image, Word, Excel, CSV or text."
    except Exception as exc:  # corrupt or password-protected files
        log.warning("could not read %s: %s", name, exc)
        doc["note"] = f"Couldn't read this file ({type(exc).__name__})."
    doc["text"] = doc["text"].strip()
    if len(doc["text"]) > MAX_TEXT_CHARS:
        doc["text"] = doc["text"][:MAX_TEXT_CHARS]
        doc["note"] = (doc["note"] + " Only the first part of this long file was read.").strip()
    if not doc["text"] and not doc["note"]:
        doc["note"] = "No readable text found in this file."
    return doc


def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "cp1252"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _read_pdf(data: bytes, llm) -> tuple[str, str]:
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted and not reader.decrypt(""):
        return "", "This PDF is password-protected. Remove the password and upload it again."
    pages = [f"[Page {i}]\n{(p.extract_text() or '').strip()}" for i, p in enumerate(reader.pages, 1)]
    text = "\n\n".join(pages)
    if len(re.sub(r"\[Page \d+\]|\s", "", text)) >= 20 * max(1, len(reader.pages)):  # ~20+ chars per page
        return text, ""
    # little or no text layer: probably a scan, so transcribe the page images
    scans = []
    for i, page in enumerate(reader.pages[:MAX_SCANNED_PAGES], 1):
        for img in list(page.images)[:1]:
            t, _ = _read_image(img.data, llm)
            if t:
                scans.append(f"[Page {i}, scanned]\n{t}")
    if scans:
        return "\n\n".join(scans), "Scanned PDF: text was read from page images, so check important numbers."
    return "", ("This PDF looks scanned (no text layer). Configure an LLM provider in .env to read scans, "
                "or upload a text-based PDF.")


def _read_image(data: bytes, llm) -> tuple[str, str]:
    if llm is None or getattr(llm, "is_mock", True):
        return "", "Images can be read once an LLM provider (OpenAI, Gemini or Claude) is set in .env."
    from PIL import Image
    img = Image.open(io.BytesIO(data))
    img.seek(0)  # first frame of animated GIFs
    img = img.convert("RGB")
    img.thumbnail((MAX_IMAGE_SIDE, MAX_IMAGE_SIDE))
    buf = io.BytesIO()
    img.save(buf, format="JPEG", quality=85)
    text = llm.describe_image(buf.getvalue(), "image/jpeg", IMAGE_PROMPT)
    return (text, "") if text else ("", "The image couldn't be read by the LLM. Try a clearer image.")


def _read_docx(data: bytes) -> str:
    import docx
    d = docx.Document(io.BytesIO(data))
    parts = [p.text for p in d.paragraphs if p.text.strip()]
    for t in d.tables:
        parts.append("\n".join(" | ".join(c.text.strip() for c in row.cells) for row in t.rows))
    return "\n".join(parts)


def _read_xlsx(data: bytes) -> str:
    import pandas as pd
    sheets = pd.read_excel(io.BytesIO(data), sheet_name=None, engine="openpyxl")
    return "\n\n".join(f"[Sheet {name}]\n{df.dropna(how='all').to_csv(index=False)}" for name, df in sheets.items())


# ---- prompt context ---------------------------------------------------------
def _pieces(text: str, size: int = 900) -> list[str]:
    """Split into ~size-char pieces on paragraph/line boundaries."""
    out, buf = [], ""
    for line in text.splitlines(keepends=True):
        if buf and len(buf) + len(line) > size:
            out.append(buf)
            buf = ""
        buf += line if len(line) <= size else line[:size]
    if buf.strip():
        out.append(buf)
    return out


def select_passages(docs: list[dict], query: str, budget: int = 6000) -> list[tuple[str, str]]:
    """Pick (file name, passage) pairs that fit `budget` characters.

    Small files go in whole. For larger ones, each file's opening is kept (so
    "summarize this" works) and the rest of the budget goes to the passages
    sharing the most words with the question, in their original order.
    """
    readable = [d for d in docs if d.get("text")]
    if sum(len(d["text"]) for d in readable) <= budget:
        return [(d["name"], d["text"]) for d in readable]
    q = set(tokenize(query))
    cands = []  # (priority, doc index, piece index, name, text)
    for di, d in enumerate(readable):
        for pi, piece in enumerate(_pieces(d["text"])):
            overlap = len(q & set(tokenize(piece)))
            cands.append((1e9 if pi == 0 else overlap, di, pi, d["name"], piece))
    chosen, used = [], 0
    for c in sorted(cands, key=lambda c: (-c[0], c[1], c[2])):
        if used + len(c[4]) <= budget:
            chosen.append(c)
            used += len(c[4])
    return [(c[3], c[4]) for c in sorted(chosen, key=lambda c: (c[1], c[2]))]


def format_documents(docs: list[dict], query: str, budget: int = 6000) -> str:
    return "\n\n".join(f"--- {name} ---\n{text.strip()}" for name, text in select_passages(docs, query, budget))
